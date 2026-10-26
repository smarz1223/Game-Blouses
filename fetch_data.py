"""
Game...Blouses data pipeline.

Downloads the two published Google Sheets workbooks (season stats + all-time history),
reads the committed draft workbook, reconciles player points to weekly scores, computes
analytics, and writes data/gameblouses_data.json. build_site.py turns that JSON into index.html.

Local testing: set STATS_FILE / HISTORY_FILE / DRAFT_FILE to local .xlsx paths.
"""
import io
import json
import math
import os
import re
import sys
import unicodedata
import urllib.request
from collections import defaultdict
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

# ----------------------------------------------------------------- CONFIG
STATS_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vRQkskVrv12zIeiOl1wxcQIqM3JkLcAEAuypUW5FKehxWKnqkrpYlbW5MZ4XXh-y4Myk6WVVYz-76sD/pub?output=xlsx")
HISTORY_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vRZGBGgjE17sJx3-i_qhfBTLLkRkUrPoqu39Uzzj8tNKoyUevH0gAy0uB-aki_BoOvzup0uJhQJQrKj/pub?output=xlsx")
DRAFT_PATH = "draft_2026.xlsx"
OUT_PATH = os.path.join("data", "gameblouses_data.json")

SEASON = 2026                 # season start year (2026-27)
REG_WEEKS = 20                # regular season weeks 1-20
PLAYOFF_TEAMS = 4             # playoffs weeks 21-22
PLAYOFF_WEEKS = [21, 22]
FEE = 75
PAYOUTS = [                   # (label, amount, how it is decided)
    ("Champion", 375, "Wins the Yahoo playoff final (week 22)"),
    ("Runner-up", 75, "Loses the Yahoo playoff final (week 22)"),
    ("Regular season", 300, "Best regular season record; Points For breaks ties"),
]
DEFAULT_MODS = {"PTS": 1.0, "REB": 1.2, "AST": 1.5, "ST": 3.0, "BLK": 3.0, "TO": -1.0}
STAT_KEYS = ["PTS", "REB", "AST", "ST", "BLK", "TO"]
STAT_LABELS = {"PTS": "Points", "REB": "Rebounds", "AST": "Assists",
               "ST": "Steals", "BLK": "Blocks", "TO": "Turnovers"}
# Position group = LAST listed eligibility (G,F -> Forward; F,C -> Center). Display order C, F, G.
GROUP_NAME = {"C": "Center", "F": "Forward", "G": "Guard"}
GROUP_ORDER = ["Center", "Forward", "Guard"]

flags = []


def flag(msg):
    flags.append(msg)
    print("FLAG:", msg)


# ----------------------------------------------------------------- LOADING
def load_wb(env_name, url, required=True):
    path = os.environ.get(env_name)
    try:
        if path:
            return load_workbook(path, data_only=True)
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=90) as r:
            return load_workbook(io.BytesIO(r.read()), data_only=True)
    except Exception as e:
        if required:
            raise
        flag(f"Could not load {env_name}: {e}")
        return None


def num(v):
    if v is None or v == "" or v == "-":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def r1(x):
    return round(x + 0.0, 1)


def r2(x):
    return round(x + 0.0, 2)


def owner_name(v):
    return re.sub(r"\s+", " ", str(v).strip()).upper()


# ----------------------------------------------------------------- MODIFIERS
def read_modifiers(wb):
    mods = dict(DEFAULT_MODS)
    if "SCORING MODIFIERS" not in wb.sheetnames:
        flag("SCORING MODIFIERS tab missing; using defaults")
        return mods
    for row in wb["SCORING MODIFIERS"].iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        m = re.search(r"\(([A-Z]+)\)", str(row[0]))
        if m and m.group(1) in mods and row[1] is not None:
            mods[m.group(1)] = float(row[1])
    return mods


# ----------------------------------------------------------------- WEEKLY SCORES
def read_weekly(wb):
    ws = wb["Weekly Scores"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and "Week" in [str(c).strip() if c else "" for c in r])
    hdr = [str(c).strip() if c else "" for c in rows[hdr_i]]
    ci = {h: i for i, h in enumerate(hdr)}
    owners, entries = [], []
    for r in rows[hdr_i + 1:]:
        if not r or r[ci["Week"]] in (None, "") or r[ci["Team"]] in (None, ""):
            continue
        o = owner_name(r[ci["Team"]])
        if o not in owners:
            owners.append(o)
        pf, pa = r[ci["Points For"]], r[ci["Points Against"]]
        if pf in (None, "") or pa in (None, ""):
            continue
        entries.append({"week": int(num(r[ci["Week"]])), "owner": o, "pf": r2(num(pf)), "pa": r2(num(pa))})
    owners = list(dict.fromkeys(owners))
    return owners, entries


def build_matchups(entries, owners):
    by_week = defaultdict(list)
    for e in entries:
        by_week[e["week"]].append(e)
    weeks, results = [], []
    for wk in sorted(by_week):
        es = by_week[wk]
        if len(es) != len(owners):
            flag(f"Week {wk}: {len(es)} of {len(owners)} scores entered; week skipped until complete")
            continue
        weeks.append(wk)
        used, games = set(), []
        for a in es:
            if a["owner"] in used:
                continue
            opp = [b for b in es if b["owner"] not in used and b["owner"] != a["owner"]
                   and abs(b["pf"] - a["pa"]) < 0.01 and abs(b["pa"] - a["pf"]) < 0.01]
            if not opp:
                flag(f"Week {wk}: no mirrored opponent found for {a['owner']}")
                continue
            b = opp[0]
            used.update([a["owner"], b["owner"]])
            games.append((a, b))
        for a, b in games:
            for x, y in ((a, b), (b, a)):
                res = "W" if x["pf"] > y["pf"] else ("L" if x["pf"] < y["pf"] else "T")
                results.append({"week": wk, "owner": x["owner"], "opp": y["owner"],
                                "pf": x["pf"], "pa": y["pf"], "res": res,
                                "playoff": wk > REG_WEEKS})
    return weeks, results


# ----------------------------------------------------------------- PLAYERS
NAME_RE = re.compile(
    r"^(?P<name>.*?)(?P<status>INJ|SUSP|GTD|DTD|NA|O|D)?"
    r"(?:No new player Notes|New Player Note|Player Note)"
    r"(?P<team>[A-Z]{2,4}) - (?P<pos>[A-Z,]+)", re.S)


def parse_player(raw):
    s = str(raw).strip()
    m = NAME_RE.match(s)
    if not m:
        first = s.split("\n")[0]
        m2 = re.match(r"^(.*?)\s*\(?([A-Z]{2,4}) - ([A-Z,]+)\)?", first)
        if m2:
            return m2.group(1).strip(), None, m2.group(2), m2.group(3).split(",")
        return first, None, "", []
    return m.group("name").strip(), m.group("status"), m.group("team"), m.group("pos").split(",")


def pos_group(positions):
    for p in reversed(positions):
        if p in GROUP_NAME:
            return GROUP_NAME[p]
    return "Forward"


def fantasy_points(stats, mods):
    return sum(stats.get(k, 0) * mods[k] for k in STAT_KEYS)


def read_players(wb, mods, owners):
    ws = wb["Consolidated Data"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and "Name" in [str(c).strip() if c else "" for c in r])
    hdr = [str(c).strip() if c else "" for c in rows[hdr_i]]
    name_i = hdr.index("Name")
    gp_i = next(i for i, h in enumerate(hdr) if h.replace("*", "") == "GP")
    stat_i = {k: hdr.index(k) for k in STAT_KEYS}
    players, log_totals = [], {}
    for r in rows[hdr_i + 1:]:
        if not r or not r[0] or r[name_i] in (None, ""):
            continue
        o = owner_name(r[0])
        stats = {k: num(r[i]) for k, i in stat_i.items()}
        if str(r[name_i]).strip() == "Totals":
            log_totals[o] = r2(fantasy_points(stats, mods))
            continue
        name, status, nba, pos = parse_player(r[name_i])
        fp = fantasy_points(stats, mods)
        gp = int(num(r[gp_i]))
        players.append({"owner": o, "player": name, "key": name_key(name), "nba": nba,
                        "positions": pos, "pos": ",".join(pos), "group": pos_group(pos),
                        "status": status or "", "gp": gp, "stats": {k: int(v) for k, v in stats.items()},
                        "fp": r1(fp), "fpg": r1(fp / gp) if gp else 0.0,
                        "cat_fp": {k: r1(stats[k] * mods[k]) for k in STAT_KEYS}})
    unknown = sorted({p["owner"] for p in players} - set(owners))
    if unknown:
        flag(f"Game log owners not in Weekly Scores: {', '.join(unknown)}")
    return players, log_totals


def name_key(name):
    s = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\b(jr|sr|ii|iii|iv)\b\.?", "", s)
    return re.sub(r"[^a-z]", "", s)


def short_key(name):
    """First initial + last name, used to match 'N. Alexander-Walker' to 'Nickeil Alexander-Walker'."""
    s = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\b(jr|sr|ii|iii|iv)\b\.?", "", s).strip()
    parts = s.split()
    if len(parts) < 2:
        return re.sub(r"[^a-z]", "", s)
    return parts[0][0] + re.sub(r"[^a-z]", "", " ".join(parts[1:]))


# ----------------------------------------------------------------- STANDINGS / ANALYTICS
def phi(z):
    return 0.5 * (1 + math.erf(z / math.sqrt(2)))


def build_standings(owners, weeks, results):
    reg = [r for r in results if not r["playoff"]]
    reg_weeks = [w for w in weeks if w <= REG_WEEKS]
    st = {o: {"owner": o, "w": 0, "l": 0, "t": 0, "pf": 0.0, "pa": 0.0, "games": 0,
              "apw": 0, "apl": 0, "apt": 0, "high": 0, "low": 0, "results": []} for o in owners}
    for r in sorted(reg, key=lambda x: x["week"]):
        s = st[r["owner"]]
        s["games"] += 1
        s["pf"] += r["pf"]
        s["pa"] += r["pa"]
        s[{"W": "w", "L": "l", "T": "t"}[r["res"]]] += 1
        s["results"].append(r["res"])
    weekly_rank = {o: [] for o in owners}
    for wk in reg_weeks:
        wr = [r for r in reg if r["week"] == wk]
        scores = sorted((r["pf"] for r in wr), reverse=True)
        for r in wr:
            s = st[r["owner"]]
            s["apw"] += sum(1 for x in wr if x["owner"] != r["owner"] and x["pf"] < r["pf"])
            s["apl"] += sum(1 for x in wr if x["owner"] != r["owner"] and x["pf"] > r["pf"])
            s["apt"] += sum(1 for x in wr if x["owner"] != r["owner"] and x["pf"] == r["pf"])
            if r["pf"] == scores[0]:
                s["high"] += 1
            if r["pf"] == scores[-1]:
                s["low"] += 1
            weekly_rank[r["owner"]].append({"week": wk, "rank": scores.index(r["pf"]) + 1})
    avgs = [s["pf"] / s["games"] for s in st.values() if s["games"]]
    mean = sum(avgs) / len(avgs) if avgs else 0
    sd = math.sqrt(sum((a - mean) ** 2 for a in avgs) / len(avgs)) if avgs else 0
    rows = []
    for o, s in st.items():
        g = s["games"]
        wpct = (s["w"] + 0.5 * s["t"]) / g if g else 0
        apg = s["apw"] + s["apl"] + s["apt"]
        appct = (s["apw"] + 0.5 * s["apt"]) / apg if apg else 0
        avg = s["pf"] / g if g else 0
        streak = ""
        if s["results"]:
            last = s["results"][-1]
            n = 0
            for x in reversed(s["results"]):
                if x != last:
                    break
                n += 1
            streak = f"{last}{n}"
        rows.append({"owner": o, "w": s["w"], "l": s["l"], "t": s["t"], "games": g,
                     "wpct": round(wpct, 3), "pf": r1(s["pf"]), "pa": r1(s["pa"]),
                     "diff": r1(s["pf"] - s["pa"]), "avg": r1(avg),
                     "ap_w": s["apw"], "ap_l": s["apl"], "ap_t": s["apt"], "ap_pct": round(appct, 3),
                     "xw": r1(appct * g), "luck": round(50 + 50 * (wpct - appct), 1) if g else 50.0,
                     "perf": round(100 * phi((avg - mean) / sd), 1) if g and sd else 50.0,
                     "high": s["high"], "low": s["low"], "streak": streak,
                     "form": s["results"][-5:]})
    rows.sort(key=lambda x: (-x["wpct"], -x["pf"]))
    for i, r in enumerate(rows, 1):
        r["seed"] = i
        r["gb"] = 0.0
    if rows:
        top = rows[0]
        for r in rows:
            r["gb"] = round(((top["w"] - r["w"]) + (r["l"] - top["l"])) / 2, 1)
    return rows, weekly_rank


def build_weekly(owners, weeks, results):
    out = []
    for wk in weeks:
        wr = [r for r in results if r["week"] == wk]
        scores = sorted((r["pf"] for r in wr), reverse=True)
        median = (scores[len(scores) // 2 - 1] + scores[len(scores) // 2]) / 2 if len(scores) % 2 == 0 else scores[len(scores) // 2]
        seen, games = set(), []
        for r in sorted(wr, key=lambda x: -x["pf"]):
            if r["owner"] in seen:
                continue
            seen.update([r["owner"], r["opp"]])
            games.append({"a": r["owner"], "a_pts": r["pf"], "b": r["opp"], "b_pts": r["pa"],
                          "margin": r1(r["pf"] - r["pa"]), "res": r["res"]})
        out.append({"week": wk, "playoff": wk > REG_WEEKS, "games": games,
                    "scores": {r["owner"]: r["pf"] for r in wr},
                    "high": max(wr, key=lambda x: x["pf"])["owner"],
                    "low": min(wr, key=lambda x: x["pf"])["owner"],
                    "avg": r1(sum(scores) / len(scores)), "median": r1(median)})
    return out


def build_h2h(owners, results):
    m = {a: {b: {"w": 0, "l": 0, "t": 0} for b in owners if b != a} for a in owners}
    for r in results:
        if r["playoff"]:
            continue
        m[r["owner"]][r["opp"]][{"W": "w", "L": "l", "T": "t"}[r["res"]]] += 1
    return m


def build_categories(players, owners, mods):
    teams = []
    for o in owners:
        mine = [p for p in players if p["owner"] == o]
        tot = {k: sum(p["stats"][k] for p in mine) for k in STAT_KEYS}
        teams.append({"owner": o, "totals": tot, "gp": sum(p["gp"] for p in mine),
                      "cat_fp": {k: r1(tot[k] * mods[k]) for k in STAT_KEYS},
                      "fp": r1(sum(tot[k] * mods[k] for k in STAT_KEYS))})
    n = len(teams) or 1
    avg = {k: r1(sum(t["cat_fp"][k] for t in teams) / n) for k in STAT_KEYS}
    for t in teams:
        t["vs_avg"] = {k: r1(t["cat_fp"][k] - avg[k]) for k in STAT_KEYS}
        t["per_game"] = {k: round(t["totals"][k] / t["gp"], 2) if t["gp"] else 0 for k in STAT_KEYS}
    leaders = {}
    for k in STAT_KEYS:
        ranked = sorted(players, key=lambda p: (p["stats"][k], -p["gp"]), reverse=True)
        leaders[k] = [{"player": p["player"], "owner": p["owner"], "val": p["stats"][k], "gp": p["gp"]}
                      for p in ranked if p["gp"]][:10]
    return {"teams": teams, "avg": avg, "leaders": leaders}


def build_positions(players, owners):
    by_owner = []
    for o in owners:
        mine = [p for p in players if p["owner"] == o]
        g = {grp: r1(sum(p["fp"] for p in mine if p["group"] == grp)) for grp in GROUP_ORDER}
        c = {grp: sum(1 for p in mine if p["group"] == grp and p["gp"]) for grp in GROUP_ORDER}
        tot = sum(g.values()) or 1
        by_owner.append({"owner": o, "fp": g, "count": c, "share": {k: round(v / tot, 3) for k, v in g.items()}})
    n = len(by_owner) or 1
    avg = {grp: r1(sum(b["fp"][grp] for b in by_owner) / n) for grp in GROUP_ORDER}
    top = {}
    for grp in GROUP_ORDER:
        ps = sorted([p for p in players if p["group"] == grp and p["gp"]], key=lambda p: -p["fp"])
        top[grp] = [{"player": p["player"], "owner": p["owner"], "nba": p["nba"], "pos": p["pos"],
                     "gp": p["gp"], "fp": p["fp"], "fpg": p["fpg"]} for p in ps[:15]]
    return {"owners": by_owner, "avg": avg, "top": top, "order": GROUP_ORDER}


# ----------------------------------------------------------------- DRAFT
def read_draft(owners):
    path = os.environ.get("DRAFT_FILE", DRAFT_PATH)
    if not os.path.exists(path):
        flag(f"Draft file {path} not found; Draft page hidden")
        return []
    ws = load_workbook(path, data_only=True)["Draft"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() for c in rows[0]]
    ci = {h: i for i, h in enumerate(hdr)}
    picks = []
    for r in rows[1:]:
        if not r or r[ci["Pick"]] is None:
            continue
        o = owner_name(r[ci["Owner"]])
        if o not in owners:
            flag(f"Draft owner {o} not in Weekly Scores")
        pos = str(r[ci["Positions"]]).split(",")
        picks.append({"pick": int(r[ci["Pick"]]), "round": int(r[ci["Round"]]), "owner": o,
                      "player": str(r[ci["Player"]]).strip(), "nba": str(r[ci["NBA Team"]]).strip(),
                      "pos": ",".join(pos), "group": pos_group(pos)})
    return sorted(picks, key=lambda p: p["pick"])


def build_draft(picks, players, owners):
    # one record per unique player across all owners
    uniq = {}
    for p in players:
        u = uniq.setdefault(p["key"], {"player": p["player"], "key": p["key"], "skey": short_key(p["player"]),
                                       "group": p["group"], "nba": p["nba"], "pos": p["pos"],
                                       "fp": 0.0, "gp": 0, "by_owner": {}})
        u["fp"] += p["fp"]
        u["gp"] += p["gp"]
        u["by_owner"][p["owner"]] = u["by_owner"].get(p["owner"], 0) + p["fp"]
    skeys = defaultdict(list)
    for u in uniq.values():
        skeys[u["skey"]].append(u)

    def find(pk):
        k = name_key(pk["player"])
        if k in uniq:
            return uniq[k]
        cands = skeys.get(short_key(pk["player"]), [])
        if len(cands) == 1:
            return cands[0]
        return None

    drafted_keys = set()
    for pk in picks:
        u = find(pk)
        pk["fp"] = r1(u["fp"]) if u else 0.0
        pk["gp"] = u["gp"] if u else 0
        pk["fp_for_owner"] = r1(u["by_owner"].get(pk["owner"], 0)) if u else 0.0
        pk["now_with"] = sorted(u["by_owner"], key=lambda o: -u["by_owner"][o]) if u else []
        pk["started"] = bool(u and u["gp"])
        if u:
            drafted_keys.add(u["key"])
    # positional ranks: draft order within group vs current FP rank within group (all players who played)
    for grp in GROUP_ORDER:
        g_picks = [p for p in picks if p["group"] == grp]
        for i, p in enumerate(sorted(g_picks, key=lambda x: x["pick"]), 1):
            p["draft_pos_rank"] = i
        pool = sorted([u for u in uniq.values() if u["group"] == grp and u["gp"]], key=lambda u: -u["fp"])
        rank_of = {u["key"]: i for i, u in enumerate(pool, 1)}
        last = len(pool) + 1
        for p in g_picks:
            u = find(p)
            p["cur_pos_rank"] = rank_of.get(u["key"], last) if u and u["gp"] else last
            p["pool_size"] = len(pool)
            p["value"] = p["draft_pos_rank"] - p["cur_pos_rank"]
    # overall rank
    pool_all = sorted([u for u in uniq.values() if u["gp"]], key=lambda u: -u["fp"])
    ov = {u["key"]: i for i, u in enumerate(pool_all, 1)}
    for p in picks:
        u = find(p)
        p["cur_ovr_rank"] = ov.get(u["key"], len(pool_all) + 1) if u else len(pool_all) + 1
    unmatched = [p["player"] for p in picks if not p["started"]]
    if unmatched:
        flag(f"Drafted players with no games logged yet: {', '.join(unmatched)}")
    # undrafted contributors
    adds = []
    for p in players:
        if p["key"] not in drafted_keys and p["gp"]:
            adds.append({"player": p["player"], "owner": p["owner"], "nba": p["nba"], "pos": p["pos"],
                         "group": p["group"], "gp": p["gp"], "fp": p["fp"]})
    adds.sort(key=lambda x: -x["fp"])
    summary = []
    for o in owners:
        mine = [p for p in picks if p["owner"] == o]
        own_fp = sum(p["fp"] for p in players if p["owner"] == o)
        draft_fp = sum(p["fp_for_owner"] for p in mine)
        summary.append({"owner": o, "picks": len(mine), "draft_fp": r1(draft_fp),
                        "other_fp": r1(own_fp - draft_fp), "total_fp": r1(own_fp),
                        "producing": sum(1 for p in mine if p["fp_for_owner"] > 0),
                        "hits": sum(1 for p in mine if p["value"] > 0),
                        "misses": sum(1 for p in mine if p["value"] < 0),
                        "avg_value": round(sum(p["value"] for p in mine) / len(mine), 1) if mine else 0})
    summary.sort(key=lambda x: -x["draft_fp"])
    order = [p["owner"] for p in picks if p["round"] == 1]
    return {"picks": picks, "adds": adds[:25], "summary": summary,
            "rounds": max(p["round"] for p in picks), "order": order}


# ----------------------------------------------------------------- HISTORY
def season_label(y):
    return f"{y}-{(y + 1) % 100:02d}"


def build_history(wb):
    if wb is None or "Data" not in wb.sheetnames:
        flag("History workbook missing Data tab")
        return None
    rows = list(wb["Data"].iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "Years")
    hdr = [str(c).strip() if c else "" for c in rows[hdr_i]]
    ci = {h: i for i, h in enumerate(hdr)}
    recs = []
    for r in rows[hdr_i + 1:]:
        if not r or r[ci["Years"]] in (None, "") or not r[ci["Manager"]]:
            continue
        y = int(num(r[ci["Years"]]))
        recs.append({"year": y, "season": season_label(y), "owner": owner_name(r[ci["Manager"]]),
                     "active": str(r[ci["Status"]]).strip().lower() == "active",
                     "first": int(num(r[ci["1st"]])), "second": int(num(r[ci["2nd"]])), "third": int(num(r[ci["3rd"]])),
                     "w": int(num(r[ci["Wins"]])), "l": int(num(r[ci["Losses"]])), "t": int(num(r[ci["Ties"]])),
                     "games": int(num(r[ci["Games"]])), "playoffs": int(num(r[ci["Playoffs"]]))})
    years = sorted({r["year"] for r in recs})
    for r in recs:
        r["wpct"] = round((r["w"] + 0.5 * r["t"]) / r["games"], 3) if r["games"] else 0
    # season finish by wins (rank within season)
    for y in years:
        ys = sorted([r for r in recs if r["year"] == y], key=lambda r: -r["wpct"])
        for r in ys:
            r["reg_rank"] = 1 + sum(1 for x in ys if x["wpct"] > r["wpct"])
    mgrs = {}
    for r in recs:
        m = mgrs.setdefault(r["owner"], {"owner": r["owner"], "active": r["active"], "seasons": 0, "w": 0, "l": 0,
                                         "t": 0, "games": 0, "titles": 0, "seconds": 0, "thirds": 0, "playoffs": 0,
                                         "best_reg": 0, "years": []})
        m["active"] = m["active"] or r["active"]
        m["seasons"] += 1
        for k in ("w", "l", "t", "games", "playoffs"):
            m[k] += r[k]
        m["titles"] += r["first"]
        m["seconds"] += r["second"]
        m["thirds"] += r["third"]
        m["best_reg"] += 1 if r["reg_rank"] == 1 else 0
        m["years"].append(r["year"])
    for m in mgrs.values():
        m["wpct"] = round((m["w"] + 0.5 * m["t"]) / m["games"], 3) if m["games"] else 0
        m["po_pct"] = round(m["playoffs"] / m["seasons"], 3) if m["seasons"] else 0
        m["first_year"] = season_label(min(m["years"]))
        ys = sorted(m["years"])
        made = {r["year"] for r in recs if r["owner"] == m["owner"] and r["playoffs"]}
        best = cur = 0
        for y in ys:
            cur = cur + 1 if y in made else 0
            best = max(best, cur)
        m["po_streak"] = best
        drought = 0
        for y in reversed(ys):
            if y in made:
                break
            drought += 1
        m["po_drought"] = drought
        del m["years"]
    managers = sorted(mgrs.values(), key=lambda m: (not m["active"], -m["wpct"]))
    champs = []
    for y in reversed(years):
        ys = [r for r in recs if r["year"] == y]
        def who(k):
            return next((r["owner"] for r in ys if r[k]), "")
        c = next((r for r in ys if r["first"]), None)
        top_w = max(r["wpct"] for r in ys)
        champs.append({"year": y, "season": season_label(y), "champ": who("first"), "runner": who("second"),
                       "third": who("third"), "champ_rec": f"{c['w']}-{c['l']}" if c else "",
                       "reg_best": [r["owner"] for r in ys if r["wpct"] == top_w],
                       "reg_best_rec": next(f"{r['w']}-{r['l']}" for r in ys if r["wpct"] == top_w),
                       "teams": len(ys)})
    best = sorted(recs, key=lambda r: (-r["wpct"], -r["w"], -r["year"]))[:8]
    worst = sorted(recs, key=lambda r: (r["wpct"], r["w"], -r["year"]))[:8]
    slim = lambda r: {"owner": r["owner"], "season": r["season"], "rec": f"{r['w']}-{r['l']}", "wpct": r["wpct"],
                      "result": "Champion" if r["first"] else ("Runner-up" if r["second"] else ("Third" if r["third"] else ("Playoffs" if r["playoffs"] else "Missed")))}
    def top(key, label, fmt=lambda v: str(v)):
        v = max(m[key] for m in managers)
        return {"label": label, "value": fmt(v), "who": [m["owner"] for m in managers if m[key] == v]}
    awards = [
        top("titles", "Most championships"),
        top("playoffs", "Most playoff trips"),
        top("po_streak", "Longest playoff streak", lambda v: f"{v} seasons"),
        top("best_reg", "Most regular season #1 finishes"),
        top("w", "Most all-time wins"),
        {"label": "Best single season", "value": slim(best[0])["rec"],
         "who": [f"{r['owner']} ({r['season']})" for r in best if r["wpct"] == best[0]["wpct"]]},
        {"label": "Worst single season", "value": slim(worst[0])["rec"],
         "who": [f"{r['owner']} ({r['season']})" for r in worst if r["wpct"] == worst[0]["wpct"]]},
    ]
    act = [m for m in managers if m["active"]]
    if act:
        d = max(m["po_drought"] for m in act)
        awards.append({"label": "Longest active playoff drought", "value": f"{d} seasons",
                       "who": [m["owner"] for m in act if m["po_drought"] == d]})
    grid = {m["owner"]: {} for m in managers}
    for r in recs:
        grid[r["owner"]][str(r["year"])] = {"rec": f"{r['w']}-{r['l']}" + (f"-{r['t']}" if r["t"] else ""),
                                            "wpct": r["wpct"], "first": r["first"], "second": r["second"],
                                            "third": r["third"], "playoffs": r["playoffs"], "reg_rank": r["reg_rank"]}
    return {"years": years, "labels": {str(y): season_label(y) for y in years}, "managers": managers,
            "champs": champs, "best": [slim(r) for r in best], "worst": [slim(r) for r in worst],
            "awards": awards, "grid": grid}


# ----------------------------------------------------------------- RECONCILIATION
def reconcile(owners, players, log_totals, results):
    out = []
    for o in owners:
        ps = r1(sum(p["fp"] for p in players if p["owner"] == o))
        wk = r1(sum(r["pf"] for r in results if r["owner"] == o))
        lt = log_totals.get(o)
        diff = r1(ps - wk)
        out.append({"owner": o, "players": ps, "weekly": wk, "diff": diff,
                    "log_total": lt, "log_ok": lt is None or abs(lt - ps) < 0.05})
        if lt is not None and abs(lt - ps) >= 0.05:
            flag(f"{o}: player rows ({ps}) do not sum to the game log Totals row ({lt})")
    return out


# ----------------------------------------------------------------- MAIN
def main():
    swb = load_wb("STATS_FILE", STATS_URL)
    hwb = load_wb("HISTORY_FILE", HISTORY_URL, required=False)
    mods = read_modifiers(swb)
    owners, entries = read_weekly(swb)
    weeks, results = build_matchups(entries, owners)
    players, log_totals = read_players(swb, mods, owners)
    standings, weekly_rank = build_standings(owners, weeks, results)
    weekly = build_weekly(owners, weeks, results)
    recon = reconcile(owners, players, log_totals, results)
    picks = read_draft(owners)
    draft = build_draft(picks, players, owners) if picks else None
    now = datetime.now(timezone.utc).astimezone(ZoneInfo("America/New_York"))
    reg_played = len([w for w in weeks if w <= REG_WEEKS])
    data = {
        "meta": {"league": "Game...Blouses", "season": season_label(SEASON), "updated": now.strftime("%b %-d, %Y %-I:%M %p ET"),
                 "weeks_played": weeks, "reg_weeks": REG_WEEKS, "reg_played": reg_played,
                 "playoff_teams": PLAYOFF_TEAMS, "playoff_weeks": PLAYOFF_WEEKS,
                 "fee": FEE, "pot": FEE * len(owners), "teams": len(owners)},
        "owners": owners,
        "mods": mods, "stat_labels": STAT_LABELS, "stat_keys": STAT_KEYS,
        "payouts": [{"label": a, "amount": b, "how": c} for a, b, c in PAYOUTS],
        "standings": standings, "weekly_rank": weekly_rank, "weekly": weekly,
        "results": results, "h2h": build_h2h(owners, results),
        "categories": build_categories(players, owners, mods),
        "positions": build_positions(players, owners),
        "players": sorted(players, key=lambda p: -p["fp"]),
        "draft": draft, "history": build_history(hwb),
        "recon": recon, "flags": flags,
    }
    os.makedirs("data", exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"Wrote {OUT_PATH}: {len(owners)} owners, weeks {weeks}, {len(players)} player rows, {len(flags)} flags")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print("FATAL:", e)
        raise
